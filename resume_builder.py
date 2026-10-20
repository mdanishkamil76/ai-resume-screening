from io import BytesIO

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


def build_resume_docx(data):
    document = Document()

    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    document.styles["Normal"].font.name = "Arial"
    document.styles["Normal"].font.size = Pt(10)

    title = data.get("resume_title") or "Professional Resume"

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = paragraph.add_run(title)
    run.bold = True
    run.font.size = Pt(16)

    target = data.get("target_role", "")
    if target:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run(f"Target Role: {target}").italic = True

    def add_heading(title):
        if title:
            document.add_heading(title, level=1)

    def add_bullets(items):
        for item in items or []:
            if isinstance(item, str) and item.strip():
                document.add_paragraph(
                    item.strip(),
                    style="List Bullet"
                )

    # Professional summary
    summary = data.get("professional_summary", "")
    if summary:
        add_heading("Professional Summary")
        document.add_paragraph(summary)

    # Core skills
    skills = data.get("core_skills", [])
    if skills:
        add_heading("Core Skills")
        document.add_paragraph(
            " | ".join(
                item for item in skills
                if isinstance(item, str)
            )
        )

    # Work experience
    experience = data.get("experience", [])
    if experience:
        add_heading("Professional Experience")

        for job in experience:
            heading = " — ".join(
                value for value in [
                    job.get("job_title", ""),
                    job.get("employer", "")
                ]
                if value
            )

            if heading:
                paragraph = document.add_paragraph()
                paragraph.add_run(heading).bold = True

            metadata = " | ".join(
                value for value in [
                    job.get("dates", ""),
                    job.get("location", "")
                ]
                if value
            )

            if metadata:
                document.add_paragraph(metadata)

            add_bullets(job.get("bullets", []))

    # Education
    education = data.get("education", [])
    if education:
        add_heading("Education")

        for item in education:
            heading = " — ".join(
                value for value in [
                    item.get("qualification", ""),
                    item.get("institution", ""),
                    item.get("dates", "")
                ]
                if value
            )

            if heading:
                document.add_paragraph(heading)

            if item.get("details"):
                document.add_paragraph(item["details"])

    # Certifications
    certifications = data.get("certifications", [])
    if certifications:
        add_heading("Certifications")
        add_bullets(certifications)

    # Projects
    projects = data.get("projects", [])
    if projects:
        add_heading("Projects")

        for project in projects:
            name = project.get("name", "")

            if name:
                paragraph = document.add_paragraph()
                paragraph.add_run(name).bold = True

            add_bullets(project.get("details", []))

    # Additional sections
    for item in data.get("additional_sections", []):
        add_heading(item.get("heading", ""))
        add_bullets(item.get("items", []))

    # Information requiring confirmation
    if data.get("items_to_confirm"):
        add_heading("Details to Confirm")

        add_bullets(data["items_to_confirm"])

    output = BytesIO()
    document.save(output)

    return output.getvalue()
