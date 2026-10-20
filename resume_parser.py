import re
from io import BytesIO


def extract_resume_text(uploaded_file):
    filename = uploaded_file.name.lower()
    raw = uploaded_file.getvalue()

    if filename.endswith(".pdf"):
        import pymupdf

        document = pymupdf.open(
            stream=raw,
            filetype="pdf"
        )

        try:
            return "\n".join(
                page.get_text("text")
                for page in document
            )
        finally:
            document.close()

    elif filename.endswith(".docx"):
        from docx import Document

        document = Document(BytesIO(raw))

        parts = [
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        # Extract table content as well.
        for table in document.tables:
            for row in table.rows:
                parts.append(
                    " | ".join(
                        cell.text.strip()
                        for cell in row.cells
                    )
                )

        return "\n".join(parts)

    elif filename.endswith(".txt"):
        return raw.decode("utf-8", errors="replace")

    raise ValueError("Unsupported resume file type.")


def get_file_signals(uploaded_file, resume_text):
    filename = uploaded_file.name.lower()
    raw = uploaded_file.getvalue()
    text = resume_text or ""
    lower = text.lower()

    words = re.findall(r"\b[\w.+#/-]+\b", text)

    sections = {
        "experience": any(
            term in lower
            for term in [
                "experience",
                "employment",
                "work history",
                "professional history"
            ]
        ),
        "education": any(
            term in lower
            for term in [
                "education",
                "academic",
                "qualification"
            ]
        ),
        "skills": any(
            term in lower
            for term in [
                "skills",
                "technical skills",
                "core competencies"
            ]
        ),
        "contact": (
            bool(
                re.search(
                    r"[\w.+-]+@[\w-]+\.[\w.-]+",
                    text
                )
            )
            or bool(
                re.search(
                    r"\+?\d[\d ()-]{7,}\d",
                    text
                )
            )
        )
    }

    page_count = None

    if filename.endswith(".pdf"):
        import pymupdf

        document = pymupdf.open(
            stream=raw,
            filetype="pdf"
        )

        try:
            page_count = len(document)
        finally:
            document.close()

    return {
        "file_name": uploaded_file.name,
        "file_type": filename.rsplit(".", 1)[-1],
        "file_size_bytes": len(raw),
        "extracted_word_count": len(words),
        "extracted_character_count": len(text),
        "standard_sections_detected": sections,
        "pdf_page_count": page_count,
        "limitations": [
            "Text extraction does not guarantee ATS compatibility.",
            "Scanned PDFs are not OCR-processed.",
            "Table reading order may differ between ATS platforms.",
            "Visual layout is not comprehensively inspected."
        ]
    }
